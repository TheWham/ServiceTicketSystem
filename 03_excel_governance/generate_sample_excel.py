"""
生成用于 Excel (.xlsx) 治理测试的样本文件：
1. Sheet 1: [order_sales_details] 纯明细流水表（上百行销售流水，适合走 Text2SQL 数据库计算）
2. Sheet 2: [reimbursement_policy] 业务规则对照表（职级差旅报销矩阵，适合转化为行级自包含语义 Chunk）
"""

import os
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def create_sample_excel(output_path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()

    # 样式定义
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="微软雅黑", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # ==================== Sheet 1: 纯明细表 (order_sales_details) ====================
    ws1 = wb.active
    ws1.title = "order_sales_details"

    headers1 = ["订单流水号", "销售大区", "业务季度", "客户行业类别", "合同销售额(元)", "履约交付成本(元)", "回款结算状态"]
    ws1.append(headers1)

    regions = ["华东大区", "华北大区", "华南大区", "西南大区"]
    quarters = ["Q1", "Q2", "Q3", "Q4"]
    client_types = ["政企客户", "金融科技", "智能制造", "医疗健康"]
    statuses = ["已回款", "已回款", "已回款", "分期挂账", "逾期催收"]

    random.seed(42)
    for i in range(1, 101):
        order_id = f"SO-2024-{i:04d}"
        region = random.choice(regions)
        quarter = random.choice(quarters)
        client = random.choice(client_types)
        sales = round(random.uniform(50000, 800000), 2)
        cost = round(sales * random.uniform(0.55, 0.78), 2)
        status = random.choice(statuses)
        ws1.append([order_id, region, quarter, client, sales, cost, status])

    # 渲染 Sheet 1 表头样式
    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws1.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18

    # ==================== Sheet 2: 业务规则表 (reimbursement_policy) ====================
    ws2 = wb.create_sheet(title="reimbursement_policy")
    headers2 = ["职级代码", "职级名称", "差旅交通标准", "一类城市住宿上限(元/天)", "二类城市住宿上限(元/天)", "单笔报销审批上限(万元)", "审批决策人"]
    ws2.append(headers2)

    rules_data = [
        ["P4-P5", "初中级工程师", "高铁二等座 / 经济舱", 450, 350, 1.0, "直属主管"],
        ["P6-P7", "高级/资深专家", "高铁一等座 / 经济舱", 650, 500, 5.0, "部门总监"],
        ["P8-P9", "资深总监/首席架构师", "高铁一等座 / 公务舱", 900, 700, 20.0, "业务线VP / 总裁"],
        ["M1-M3", "集团执委会高管", "高铁商务座 / 头等舱", 1500, 1200, 100.0, "集团执行总裁"],
    ]

    for row in rules_data:
        ws2.append(row)

    # 渲染 Sheet 2 样式
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

    wb.save(output_path)
    print(f"[SUCCESS] 样本 EXCEL 成功生成至: {output_path}")


if __name__ == "__main__":
    sample_file = os.path.join(os.path.dirname(__file__), "enterprise_data.xlsx")
    create_sample_excel(sample_file)
