"""
Layer 1 数据物理清洗层：clean_pipeline.py
职责：
1. 版面解析：遍历 docx 底层 XML（<w:p>/<w:tbl>）保序提取结构化块（改造自 02_docx_governance 的保序遍历）
2. 规则清洗：分隔线/空段剔除、空表判定、水印正则框架、字段表摊平、emoji 字段行解析、PII 脱敏
3. 条目切分：按 Heading1 记大类、Heading2 切知识条目，输出纯净 Markdown（落盘 output/cleaned/）
"""

import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


# ==================== 数据结构 ====================

@dataclass
class Block:
    """保序解析出的文档块（原始态的中间表示，尚未清洗）"""
    kind: str                      # h1 | h2 | h3 | para | list | table
    text: str = ""                 # 段落文本（table 类型不用此字段）
    rows: List[List[str]] = field(default_factory=list)  # 表格内容（行 -> 单元格）


@dataclass
class Entry:
    """一条清洗后的知识条目（贯穿全链路的核心数据结构：
    规则清洗产出 -> LLM 清洗替换 fields -> LLM 萃取读取内容 -> Wiki 渲染消费）"""
    code: str = ""                 # 条目代码，如 NET-001；附录为 APPENDIX-1 等
    title: str = ""                # 条目标题
    category: str = ""             # 所属大类（H1 去序号/emoji）
    keywords: List[str] = field(default_factory=list)   # 🏷 检索关键词
    fields: Dict[str, str] = field(default_factory=dict)  # 字段表摊平结果：现象描述/常见根因/...
    remark: str = ""               # 📝 备注
    body_paras: List[str] = field(default_factory=list)   # 无结构化字段条目的正文段落（如附录）
    tables: List[List[List[str]]] = field(default_factory=list)  # 保留的真表格
    is_appendix: bool = False      # 附录标记：无【代码】的 H2 章节


# ==================== 清洗规则常量 ====================

ENTRY_CODE_RE = re.compile(r"【\s*([A-Z]+-\d+)\s*】\s*(.+)")
SEPARATOR_RE = re.compile(r"^[─—\-=_*\s]{6,}$")
H1_PREFIX_RE = re.compile(r"^[一二三四五六七八九十百]+、")
EMOJI_PREFIX_RE = re.compile(r"^[\U0001F300-\U0001FAFF☀-➿⬀-⯿️⃣\s]+")

KEYWORD_LINE_RE = re.compile(r"^🏷\s*检索关键词[：:]\s*(.+)")
REMARK_LINE_RE = re.compile(r"^📝\s*备注[：:]\s*(.+)")

# 字段表左侧标签 -> 规范字段名
FIELD_LABEL_MAP = {
    "现象描述": "现象描述",
    "常见根因": "常见根因",
    "员工自助排查": "员工自助排查",
    "IT工程师进阶处理": "IT工程师进阶处理",
}
FIELD_LABEL_RE = re.compile(r"^[\U0001F300-\U0001FAFF☀-➿\s️]*(现象描述|常见根因|员工自助排查|IT工程师进阶处理)\s*$")

# R4 水印/页眉页脚正则框架（可配置，本文档无命中）
WATERMARK_PATTERNS: List[re.Pattern] = [
    # re.compile(r"内部资料[，,]?\s*请勿外传"),
]

# R8 PII 脱敏规则：类型占位符替换（保留实体类型与语义，供 RAG 检索可用）
# 覆盖：身份证号/手机号/邮箱/银行账号/IP/护照/车牌/凭证密钥/数据库连接串/私钥/姓名(上下文规则)
# 注：MAC 地址在运维排障文档中承担技术语义（绑定排查），默认不脱敏，可按需开启
PII_PATTERNS: List[Tuple[re.Pattern, str]] = [
    # --- 凭证密钥（优先级最高，先匹配避免被其他规则切碎） ---
    (re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----[\s\S]*?-----END [A-Z ]*PRIVATE KEY-----"), "[私钥已脱敏]"),
    (re.compile(r"(?i)\b(?:mysql|postgres(?:ql)?|mongodb(?:\+srv)?|redis|mssql|sqlserver|oracle)://\S+"), "[数据库连接串]"),
    (re.compile(r"(?i)(password|passwd|pwd|token|api[_-]?key|secret[_-]?key|access[_-]?key|密码|口令)(\s*[=:：]\s*)\S+"), r"\1\2[凭证密钥]"),
    # --- 直接标识符 ---
    (re.compile(r"(?<![\d.])\d{17}[\dXx](?!\d)"), "[身份证号]"),
    (re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"), "[手机号]"),
    (re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"), "[邮箱]"),
    (re.compile(r"((?:结算)?(?:账号|账户|卡号|银行卡)[号]?[：:]?\s*)[\d][\d\s]{10,26}\d"), r"\1[银行账号]"),
    (re.compile(r"(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}(?![\d.])"), "[IP地址]"),
    (re.compile(r"(?<![A-Za-z0-9])[EGPSegps]\d{8}(?![A-Za-z0-9])"), "[护照号]"),
    (re.compile(r"[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼][A-HJ-NP-Z][-·]?[A-HJ-NP-Z0-9]{4,5}[A-HJ-NP-Z0-9挂学警港澳]"), "[车牌号]"),
    # --- 姓名（上下文规则：标签后 2-4 个汉字） ---
    (re.compile(r"(联系人|负责人|姓名|经办人|签收人)([：:]\s*)[一-龥]{2,4}"), r"\1\2[姓名]"),
    # --- 设备标识 ---
    (re.compile(r"(?<!\d)8[56]\d{13}(?!\d)"), "[IMEI]"),  # IMEI 多以 86/85 开头的 15 位
]


def desensitize_text(text: str) -> Tuple[str, int]:
    """对任意文本执行全量 PII 脱敏，返回 (脱敏后文本, 命中次数)"""
    hits = 0
    for pattern, repl in PII_PATTERNS:
        text, n = pattern.subn(repl, text)
        hits += n
    return text, hits


class CleanPipeline:
    """Layer 1：docx 物理清洗引擎"""

    def __init__(self, docx_path: str):
        self.docx_path = docx_path
        self.stats: Dict[str, int] = {}

    # ---------- 1. 版面解析（保序） ----------

    def parse_blocks(self) -> List[Block]:
        """【Step 1-a】遍历 docx 底层 XML 子节点，按物理阅读顺序提取结构化块。
        关键点：不用 doc.paragraphs / doc.tables 分别遍历（那会导致表格坠落到文末），
        而是直接遍历 body 的 <w:p>/<w:tbl> 子节点，保证标题、段落、表格顺序不错位。"""
        doc = Document(self.docx_path)
        blocks: List[Block] = []
        for child in doc.element.body:
            if isinstance(child, CT_P):  # Word 段落节点
                p = Paragraph(child, doc)
                text = p.text.strip()
                style = p.style.name
                # 按样式名映射块类型（Heading 1/2/3 用于后续分章节与切条目）
                if style.startswith("Heading 1"):
                    blocks.append(Block(kind="h1", text=text))
                elif style.startswith("Heading 2"):
                    blocks.append(Block(kind="h2", text=text))
                elif style.startswith("Heading 3"):
                    blocks.append(Block(kind="h3", text=text))
                elif style.startswith("List"):
                    blocks.append(Block(kind="list", text=text))
                else:
                    blocks.append(Block(kind="para", text=text))
            elif isinstance(child, CT_Tbl):  # Word 表格节点
                table = Table(child, doc)
                # 单元格内的换行压缩为 \n（保留行结构，供后续摊平时拆成 Markdown 列表行）
                rows = [
                    [re.sub(r"\s*\n\s*", "\n", cell.text.strip()) for cell in row.cells]
                    for row in table.rows
                ]
                blocks.append(Block(kind="table", rows=rows))
        return blocks

    # ---------- 2. 规则清洗 ----------

    def _is_empty_table(self, rows: List[List[str]]) -> bool:
        return all(not cell.strip() for row in rows for cell in row)

    def _desensitize(self, text: str) -> Tuple[str, int]:
        return desensitize_text(text)

    def clean(self, blocks: List[Block]) -> List[Block]:
        """【Step 1-b】规则清洗主循环：对每块依次应用 R1-R8，返回干净块列表。
        所有命中数计入 self.stats，供入口打印对照基线。"""
        stats = {
            "raw_blocks": len(blocks),
            "R1_分隔线剔除": 0,
            "R2_空段剔除": 0,
            "R3_空表剔除": 0,
            "R4_水印命中": 0,
            "R8_PII脱敏": 0,
        }
        cleaned: List[Block] = []
        for b in blocks:
            if b.kind == "table":
                # R3 空表判定：所有单元格为空的占位表直接剔除（本文档无此类表）
                if self._is_empty_table(b.rows):
                    stats["R3_空表剔除"] += 1
                    continue
                # R8 对表格逐单元格脱敏
                new_rows = []
                for row in b.rows:
                    new_row = []
                    for cell in row:
                        cell, n = self._desensitize(cell)
                        stats["R8_PII脱敏"] += n
                        new_row.append(cell)
                    new_rows.append(new_row)
                cleaned.append(Block(kind="table", rows=new_rows))
                continue

            text = b.text
            # R2 空段剔除
            if not text:
                stats["R2_空段剔除"] += 1
                continue
            if b.kind in ("para", "list"):
                # R1 分隔线剔除（───── 这类装饰线）
                if SEPARATOR_RE.match(text):
                    stats["R1_分隔线剔除"] += 1
                    continue
                # R4 水印/页眉正则框架（WATERMARK_PATTERNS 可配置，本文档无命中）
                if any(p.search(text) for p in WATERMARK_PATTERNS):
                    stats["R4_水印命中"] += 1
                    continue
                # R8 PII 脱敏（类型占位符，保 RAG 语义）
                text, n = self._desensitize(text)
                stats["R8_PII脱敏"] += n
            cleaned.append(Block(kind=b.kind, text=text, rows=b.rows))

        stats["clean_blocks"] = len(cleaned)
        self.stats = stats
        return cleaned

    # ---------- 3. 条目切分与字段摊平 ----------

    @staticmethod
    def _clean_category(h1_text: str) -> str:
        text = H1_PREFIX_RE.sub("", h1_text.strip())
        text = EMOJI_PREFIX_RE.sub("", text).strip()
        return text

    def split_entries(self, blocks: List[Block]) -> List[Entry]:
        """【Step 1-c】条目切分：H1 记分类、H2 切条目，字段表摊平。
        条目签名（本文档固定结构）：H2【CODE】标题 -> 🏷关键词段 -> 2列字段表 -> 可选📝备注段"""
        entries: List[Entry] = []
        category = ""
        current: Optional[Entry] = None
        appendix_seq = 0
        stats_flatten = 0

        def flush():
            if current is not None:
                entries.append(current)

        for b in blocks:
            if b.kind == "h1":
                # 遇到大类章节标题：flush 上一条目，更新当前分类
                flush()
                current = None
                category = self._clean_category(b.text)
                continue

            if b.kind == "h2":
                # 遇到条目标题：flush 上一条目，开新条目
                flush()
                m = ENTRY_CODE_RE.search(b.text)  # R7 条目代码正则【XXX-NNN】
                if m:
                    current = Entry(code=m.group(1), title=m.group(2).strip(), category=category)
                else:
                    # 附录 H2 无代码（如"员工自助排查流程"），编号为 APPENDIX-N
                    appendix_seq += 1
                    current = Entry(
                        code=f"APPENDIX-{appendix_seq}",
                        title=b.text.strip(),
                        category=category,
                        is_appendix=True,
                    )
                continue

            if current is None:
                continue  # 使用说明/分类索引等前置章节不属于任何条目，跳过

            if b.kind == "table":
                # R5 核心规则：2 列且首列是字段标签的表 -> 摊平为 Entry.fields；
                # 否则（如速查表）保留为真表格进 Entry.tables
                if b.rows and len(b.rows[0]) == 2 and FIELD_LABEL_RE.match(b.rows[0][0]):
                    for row in b.rows:
                        lm = FIELD_LABEL_RE.match(row[0])
                        if lm:
                            current.fields[FIELD_LABEL_MAP[lm.group(1)]] = row[1].strip()
                    stats_flatten += 1
                else:
                    current.tables.append(b.rows)
                continue

            text = b.text
            # R6 emoji 字段行解析：🏷 关键词 / 📝 备注
            km = KEYWORD_LINE_RE.match(text)
            if km:
                current.keywords = [k.strip() for k in re.split(r"[,，、]", km.group(1)) if k.strip()]
                continue
            rm = REMARK_LINE_RE.match(text)
            if rm:
                current.remark = rm.group(1).strip()
                continue
            current.body_paras.append(text)  # 其余段落进正文

        flush()
        self.stats["R5_字段表摊平"] = stats_flatten
        self.stats["entries"] = len(entries)
        self.stats["appendix_entries"] = sum(1 for e in entries if e.is_appendix)
        return entries

    # ---------- 4. 纯净 Markdown 输出 ----------

    @staticmethod
    def _cell_to_md_lines(text: str) -> List[str]:
        """字段单元格内的换行内容转为 Markdown 列表行"""
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
        return lines

    @staticmethod
    def _table_to_md(rows: List[List[str]]) -> str:
        header = [c.replace("\n", " ") for c in rows[0]]
        md = "| " + " | ".join(header) + " |\n"
        md += "| " + " | ".join(["---"] * len(header)) + " |\n"
        for row in rows[1:]:
            md += "| " + " | ".join(c.replace("\n", "<br>") for c in row) + " |\n"
        return md

    def entry_to_markdown(self, e: Entry) -> str:
        lines = [f"# 【{e.code}】{e.title}" if not e.is_appendix else f"# {e.title}", ""]
        if e.category:
            lines.append(f"> 所属分类：{e.category}\n")
        if e.keywords:
            lines.append("**检索关键词**：" + ", ".join(e.keywords) + "\n")

        field_titles = [("现象描述", "现象描述"), ("常见根因", "常见根因"),
                        ("员工自助排查", "员工自助排查"), ("IT工程师进阶处理", "IT工程师进阶处理")]
        for key, title in field_titles:
            if key in e.fields:
                lines.append(f"## {title}\n")
                for ln in self._cell_to_md_lines(e.fields[key]):
                    lines.append(f"- {ln}")
                lines.append("")

        for para in e.body_paras:
            lines.append(para + "\n")
        for rows in e.tables:
            lines.append(self._table_to_md(rows))
        if e.remark:
            lines.append(f"> 📝 备注：{e.remark}\n")
        return "\n".join(lines).strip() + "\n"

    # ---------- 5. 编排 ----------

    def run(self, output_dir: str) -> Tuple[List[Entry], Dict[str, int]]:
        blocks = self.parse_blocks()
        cleaned = self.clean(blocks)
        entries = self.split_entries(cleaned)

        os.makedirs(output_dir, exist_ok=True)
        merged = []
        for e in entries:
            md = self.entry_to_markdown(e)
            merged.append(md)
            with open(os.path.join(output_dir, f"{e.code}.md"), "w", encoding="utf-8") as f:
                f.write(md)
        with open(os.path.join(output_dir, "_merged_clean.md"), "w", encoding="utf-8") as f:
            f.write("\n\n---\n\n".join(merged))
        return entries, self.stats
