"""
Word (.docx) 格式专项治理引擎：
对比【朴素分离式提取（导致表格与正文倒错）】与【XML 元素保序流式提取（严格维持图文与表格上下文顺序）】
"""

from typing import List, Tuple
from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table


class DocxGovernanceEngine:
    def __init__(self, docx_path: str):
        self.docx_path = docx_path
        self.doc = Document(docx_path)

    def parse_naive_separated(self) -> str:
        """
        ❌ 朴素提取错误方式：
        分别读取 doc.paragraphs 和 doc.tables，
        导致所有表格被整体放到了文档最底部，完全割裂了段落与表格的紧密上下文！
        """
        output_lines: List[str] = ["# ❌ 朴素提取结果（上下文顺序严重错位）\n"]

        # 先遍历所有段落
        output_lines.append("## 【第一阶段：提取的所有自然段落】")
        for p in self.doc.paragraphs:
            text = p.text.strip()
            if text:
                output_lines.append(text)

        # 再遍历所有表格
        output_lines.append("\n## 【第二阶段：提取的所有表格（已被抛弃在文章末尾）】")
        for tbl_idx, table in enumerate(self.doc.tables):
            output_lines.append(f"\n[表格 #{tbl_idx + 1}]")
            for row in table.rows:
                row_str = " | ".join([cell.text.strip() for cell in row.cells])
                output_lines.append(f"| {row_str} |")

        return "\n".join(output_lines)

    def parse_stream_ordered(self) -> str:
        """
        ✅ 推荐工程策略：底层 XML 元素流式保序提取
        通过遍历 doc.element.body 的直接子节点，按物理阅读流交替提取段落与表格，
        实现真正的图文、标题与表格严格对齐。
        """
        output_lines: List[str] = ["# ✅ 保序流式提取结果（严格对齐真实阅读顺序）\n"]

        # 遍历文档 Body 下所有按自然顺序排列的 XML 子元素
        for child in self.doc.element.body:
            # 1. 如果是段落元素 (<w:p>)
            if isinstance(child, CT_P):
                p = Paragraph(child, self.doc)
                text = p.text.strip()
                if not text:
                    continue

                # 识别标题级别，转换对应 Markdown 标题语法
                if p.style.name.startswith("Heading 1"):
                    output_lines.append(f"# {text}\n")
                elif p.style.name.startswith("Heading 2"):
                    output_lines.append(f"## {text}\n")
                elif p.style.name.startswith("Heading 3"):
                    output_lines.append(f"### {text}\n")
                else:
                    output_lines.append(f"{text}\n")

            # 2. 如果是表格元素 (<w:tbl>)
            elif isinstance(child, CT_Tbl):
                table = Table(child, self.doc)
                if not table.rows:
                    continue

                # 提取表头
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                md_table = "\n| " + " | ".join(headers) + " |\n"
                md_table += "| " + " | ".join(["---"] * len(headers)) + " |\n"

                # 提取数据行
                for row in table.rows[1:]:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    md_table += "| " + " | ".join(row_cells) + " |\n"

                output_lines.append(md_table)

        return "\n".join(output_lines)
