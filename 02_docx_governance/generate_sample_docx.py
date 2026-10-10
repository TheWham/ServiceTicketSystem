"""
生成用于 Word (.docx) 治理测试的样本文件：
包含多层级标题、正文段落、穿插在段落中间的高价值业务表格，用于验证保序流式提取。
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_sample_docx(output_path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # 1. 文档大标题
    title = doc.add_heading("集团核心系统运维保障与服务等级协议 (SLA)", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 2. 序言段落
    p1 = doc.add_paragraph("本协议由集团基础设施技术保障部与各业务中心联合制定，用于规范生产故障处理链路与考核指标。")
    p2 = doc.add_paragraph("各业务系统负责人必须严格遵守以下故障定级标准与恢复预案要求。")

    # 3. 章节 1：穿插在中间的重要表格
    doc.add_heading("一、故障等级与紧急响应时限矩阵", level=2)
    doc.add_paragraph("系统突发故障发生时，各责任中心应按下表规定的时限启动应急响应，严禁瞒报或推诿：")

    # 插入复杂表格
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"

    headers = ["故障级别", "业务影响范围", "响应达标时限", "降级与容灾措施"]
    data = [
        ["P0 (灾难级)", "核心交易或结算链路中断 > 5分钟", "≤ 3 分钟", "自动触发跨机房切流，拉起 WarRoom 指挥群"],
        ["P1 (严重级)", "单一业务集群或主打业务功能不可用", "≤ 10 分钟", "熔断故障依赖，开启只读兜底降级方案"],
        ["P2 (一般级)", "非核心业务模块异常或偶发性超时", "≤ 30 分钟", "摘除异常实例，收集链路追踪 Log 工单排期"],
    ]

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = text

    for row_idx, row_data in enumerate(data):
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = text

    # 4. 紧接着表格下方的关联正文（如果脱离表格，这段正文将失去主谓语境）
    p3 = doc.add_paragraph("【特别声明】：上表中的响应时限以 APM 监控平台自动触发告警的打点时间戳为准。")

    # 5. 章节 2
    doc.add_heading("二、免责条款与不可抗力说明", level=2)
    doc.add_paragraph("由外部运营商骨干网络中断或云厂商基础机房断电引发的大规模不可抗力事件，经技术委员会评审核实后可免予扣减 SLA 绩效积分。")

    doc.save(output_path)
    print(f"[SUCCESS] 样本 DOCX 成功生成至: {output_path}")


if __name__ == "__main__":
    sample_file = os.path.join(os.path.dirname(__file__), "enterprise_sla.docx")
    create_sample_docx(sample_file)
